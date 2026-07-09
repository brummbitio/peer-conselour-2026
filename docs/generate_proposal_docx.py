from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "proposal-pengadaan-website-2026.md"
OUTPUT = ROOT / "Proposal Pengadaan Website Layanan Konseling Mahasiswa UB 2026.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def clean_inline(text: str) -> str:
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    return text.strip()


def apply_base_style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    for style_name, size in [("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12)]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True

    section = document.sections[0]
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)


def add_paragraph_with_format(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(clean_inline(text))


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = clean_inline(value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                if r_idx == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.bold = True
            if r_idx == 0:
                set_cell_shading(cell, "D9EAF7")


def build_docx() -> None:
    document = Document()
    apply_base_style(document)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == r"\newpage":
            document.add_page_break()
            i += 1
            continue

        if stripped.startswith("|"):
            table_rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip()
                if re.fullmatch(r"\|(?:\s*[-:]+\s*\|)+", raw):
                    i += 1
                    continue
                table_rows.append([cell.strip() for cell in raw.strip("|").split("|")])
                i += 1
            add_table(document, table_rows)
            continue

        if stripped.startswith("### "):
            add_paragraph_with_format(document, stripped[4:], style="Heading 3")
            i += 1
            continue

        if stripped.startswith("## "):
            add_paragraph_with_format(document, stripped[3:], style="Heading 2")
            i += 1
            continue

        if stripped.startswith("# "):
            add_paragraph_with_format(document, stripped[2:], style="Heading 1")
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            add_paragraph_with_format(document, re.sub(r"^\d+\.\s+", "", stripped), style="List Number")
            i += 1
            continue

        if stripped.startswith("- "):
            add_paragraph_with_format(document, stripped[2:], style="List Bullet")
            i += 1
            continue

        add_paragraph_with_format(document, stripped)
        i += 1

    document.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
